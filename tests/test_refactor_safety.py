"""Регрессии для сценариев, способных менять или повреждать пользовательские файлы."""
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from docx import Document

from app.core.rename_plan import resolve_rename_targets
from core.workers.rename_transaction import rename_batch
from core.workers.atomic_output import atomic_output_path
from core.workers.compression.image_optimizer import save_optimized_image
from core.workers.merge.merge_mixin import _requires_docx_composer


class RenameSafetyTests(unittest.TestCase):
    def test_swap_names_preserves_both_files(self):
        with tempfile.TemporaryDirectory() as directory:
            first, second = [Path(directory, name) for name in ('a.txt', 'b.txt')]
            first.write_text('FIRST', encoding='utf8')
            second.write_text('SECOND', encoding='utf8')
            items = [SimpleNamespace(path=str(path), folder=directory) for path in (first, second)]
            targets = resolve_rename_targets(items, ['b.txt', 'a.txt'])
            successes, errors = rename_batch([item.path for item in items], targets)
            self.assertEqual(errors, [])
            self.assertEqual(len(successes), 2)
            self.assertEqual(first.read_text(), 'SECOND')
            self.assertEqual(second.read_text(), 'FIRST')

    def test_external_target_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            first, existing = [Path(directory, name) for name in ('a.txt', 'b.txt')]
            first.write_text('FIRST')
            existing.write_text('EXISTING')
            item = SimpleNamespace(path=str(first), folder=directory)
            [target] = resolve_rename_targets([item], ['b.txt'])
            self.assertEqual(os.path.basename(target), 'b_1.txt')
            successes, errors = rename_batch([str(first)], [target])
            self.assertFalse(errors)
            self.assertEqual(existing.read_text(), 'EXISTING')
            self.assertEqual(Path(successes[0][1]).read_text(), 'FIRST')

    def test_duplicate_targets_get_distinct_names(self):
        with tempfile.TemporaryDirectory() as directory:
            items = [SimpleNamespace(path=str(Path(directory, name)), folder=directory)
                     for name in ('a.txt', 'b.txt')]
            result = resolve_rename_targets(items, ['x.txt', 'X.txt'])
            self.assertEqual(len({path.casefold() for path in result}), 2)

    def test_staging_failure_rolls_back_files(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory, name) for name in ('a.txt', 'b.txt')]
            for idx, path in enumerate(paths):
                path.write_text(str(idx))
            true_rename = os.rename
            calls = 0

            def fail_second_move(source, target):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise PermissionError('mock staging failure')
                return true_rename(source, target)

            with patch('core.workers.rename_transaction.os.rename', side_effect=fail_second_move):
                successes, errors = rename_batch(
                    [str(path) for path in paths], [str(paths[1]), str(paths[0])]
                )
            self.assertFalse(successes)
            self.assertTrue(errors)
            self.assertEqual(paths[0].read_text(), '0')
            self.assertEqual(paths[1].read_text(), '1')
            self.assertFalse(list(Path(directory).glob('.__multifora_*')))

    def test_invalid_windows_names_are_blocked(self):
        with tempfile.TemporaryDirectory() as directory:
            item = SimpleNamespace(path=str(Path(directory, 'a.txt')), folder=directory)
            for invalid in ('CON.txt', '../outside.txt', 'bad?.txt', 'end. '):
                with self.subTest(name=invalid), self.assertRaises(ValueError):
                    resolve_rename_targets([item], [invalid])


class ImageSafetyTests(unittest.TestCase):
    def test_rgba_transparency_survives_png_optimization(self):
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory, 'optimized.png'))
            image = Image.new('RGBA', (24, 24), (20, 50, 80, 0))
            image.putpixel((4, 5), (250, 10, 30, 180))
            save_optimized_image(image, output, '.png', 85)
            with Image.open(output) as result:
                self.assertEqual(result.mode, 'RGBA')
                self.assertEqual(result.getpixel((0, 0)), (20, 50, 80, 0))
                self.assertEqual(result.getpixel((4, 5)), (250, 10, 30, 180))

    def test_palette_transparency_survives_png_optimization(self):
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory, 'optimized.png'))
            image = Image.new('P', (24, 24), 0)
            palette = [0] * 768
            palette[3:6] = [255, 0, 0]
            image.putpalette(palette)
            image.info['transparency'] = 0
            image.putpixel((5, 5), 1)
            save_optimized_image(image, output, '.png', 85)
            with Image.open(output) as result:
                self.assertEqual(result.info['transparency'], 0)
                self.assertEqual(result.convert('RGBA').getpixel((0, 0))[3], 0)
                self.assertEqual(result.convert('RGBA').getpixel((5, 5))[3], 255)


class AtomicSaveTests(unittest.TestCase):
    def test_fail_does_not_publish_partial_output(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory, 'merged.pdf')
            with self.assertRaises(RuntimeError):
                with atomic_output_path(str(output)) as temporary:
                    Path(temporary).write_bytes(b'part')
                    raise RuntimeError('failed')
            self.assertFalse(output.exists())
            self.assertFalse(list(Path(directory).glob('.__multifora_*')))

    def test_success_is_published(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory, 'merged.pdf')
            with atomic_output_path(str(output)) as temporary:
                Path(temporary).write_bytes(b'%PDF content')
            self.assertEqual(output.read_bytes(), b'%PDF content')

    def test_does_not_replace_preexisting_target(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory, 'merged.pdf')
            output.write_bytes(b'original')
            with self.assertRaises(FileExistsError):
                with atomic_output_path(str(output)) as temporary:
                    Path(temporary).write_bytes(b'new')
            self.assertEqual(output.read_bytes(), b'original')


class DocxSafetyTests(unittest.TestCase):
    def test_embedded_image_requires_relationship_aware_merge(self):
        with tempfile.TemporaryDirectory() as directory:
            picture = Path(directory, 'picture.png')
            Image.new('RGB', (5, 5), 'red').save(picture)
            doc = Document()
            doc.add_picture(str(picture))
            self.assertTrue(_requires_docx_composer(doc))

    def test_plain_paragraph_needs_no_special_merger(self):
        doc = Document()
        doc.add_paragraph('Hello')
        self.assertFalse(_requires_docx_composer(doc))

class ConversionOutputSafetyTests(unittest.TestCase):
    def test_failed_pdf_to_word_removes_partial_document_and_closes_converter(self):
        from unittest.mock import Mock
        import core.workers.conversion.conversion_mixin as module
        from core.workers.conversion.conversion_mixin import ConversionMixin
        from app.core.models import FileItem

        class Worker(ConversionMixin):
            def __init__(self):
                self.conversion_output_mode = 'alongside'
                self.conversion_output_dir = ''
                self._conversion_reserved_paths = set()
            def _should_cancel(self):
                return False

        with tempfile.TemporaryDirectory() as directory:
            pdf = Path(directory, 'document.pdf')
            pdf.write_bytes(b'placeholder')
            converter = Mock()
            def partial_then_fail(destination):
                Path(destination).write_bytes(b'invalid half document')
                raise ValueError('simulated failure')
            converter.convert.side_effect = partial_then_fail
            factory = SimpleNamespace(Converter=Mock(return_value=converter))
            with patch.object(module, 'HAS_PDF_TO_WORD', True), patch.object(module, 'pdf2docx', factory):
                with self.assertRaisesRegex(Exception, 'simulated failure'):
                    Worker()._convert_pdf_to_word(FileItem(str(pdf)))
            self.assertFalse(Path(directory, 'document.docx').exists())
            self.assertFalse(list(Path(directory).glob('.__multifora_*')))
            converter.close.assert_called_once()

    def test_pdf_merge_produces_readable_pdf(self):
        import pymupdf
        from core.workers.merge.merge_mixin import MergeMixin
        from app.core.models import FileItem

        class Worker(MergeMixin):
            def __init__(self, files):
                self.files = files
                self.merge_output_format = 'pdf'
                self.merge_output_path = ''
                self.status = SimpleNamespace(emit=lambda _value: None)
                self.progress = SimpleNamespace(emit=lambda _value: None)
            def _should_cancel(self):
                return False
            def _get_unique_path(self, path):
                return path

        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory, name) for name in ('a.pdf', 'b.pdf')]
            for path in paths:
                with pymupdf.open() as pdf:
                    pdf.new_page()
                    pdf.save(str(path))
            output = Worker([FileItem(str(path)) for path in paths])._merge_files_to_target()
            with pymupdf.open(output) as merged:
                self.assertEqual(merged.page_count, 2)


if __name__ == '__main__':
    unittest.main()
