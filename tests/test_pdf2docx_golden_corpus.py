"""Small generated corpus that characterizes the production PDF backend."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core import deps

pytestmark = pytest.mark.skipif(
    not deps.HAS_PDF_TO_DOCX,
    reason="pdf2docx/PyMuPDF backend is not installed",
)


def _convert(source: Path, destination: Path) -> Document:
    converter = deps.pdf2docx.Converter(str(source))
    try:
        converter.convert(str(destination))
    finally:
        converter.close()
    return Document(destination)


def test_pdf2docx_golden_text_and_page_order(tmp_path: Path):
    import pymupdf

    source = tmp_path / "ordered-pages.pdf"
    output = tmp_path / "ordered-pages.docx"
    pdf = pymupdf.open()
    for marker in ("GOLDEN PAGE ONE", "GOLDEN PAGE TWO"):
        page = pdf.new_page()
        page.insert_text((72, 96), marker, fontsize=16)
        page.insert_text((72, 124), "Body text 12345", fontsize=11)
    pdf.save(source)
    pdf.close()

    document = _convert(source, output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "GOLDEN PAGE ONE" in text
    assert "GOLDEN PAGE TWO" in text
    assert text.index("GOLDEN PAGE ONE") < text.index("GOLDEN PAGE TWO")
    assert text.count("Body text 12345") == 2


def test_pdf2docx_golden_table_structure(tmp_path: Path):
    import pymupdf

    source = tmp_path / "table.pdf"
    output = tmp_path / "table.docx"
    pdf = pymupdf.open()
    page = pdf.new_page()
    left, top, width, height = 72, 96, 360, 80
    for x in (left, left + width / 2, left + width):
        page.draw_line((x, top), (x, top + height))
    for y in (top, top + height / 2, top + height):
        page.draw_line((left, y), (left + width, y))
    for x, y, value in (
        (left + 8, top + 25, "A1"),
        (left + width / 2 + 8, top + 25, "B1"),
        (left + 8, top + height / 2 + 25, "A2"),
        (left + width / 2 + 8, top + height / 2 + 25, "B2"),
    ):
        page.insert_text((x, y), value, fontsize=11)
    pdf.save(source)
    pdf.close()

    document = _convert(source, output)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]

    assert cells == ["A1", "B1", "A2", "B2"]


def test_pdf2docx_golden_scanned_cyrillic_image(tmp_path: Path):
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont

    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if not font_path.is_file():
        pytest.skip("Windows Arial font is unavailable")

    scan = tmp_path / "cyrillic-scan.png"
    image = Image.new("RGB", (900, 180), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (30, 45),
        "Кириллица: Привет, мир",
        fill="black",
        font=ImageFont.truetype(str(font_path), 48),
    )
    image.save(scan)

    source = tmp_path / "cyrillic-scan.pdf"
    output = tmp_path / "cyrillic-scan.docx"
    pdf = pymupdf.open()
    page = pdf.new_page(width=900, height=220)
    page.insert_image(pymupdf.Rect(0, 0, 900, 180), filename=str(scan))
    pdf.save(source)
    pdf.close()

    document = _convert(source, output)
    assert document.inline_shapes or document.part.package.parts
    assert any(part.content_type.startswith("image/") for part in document.part.package.parts)
